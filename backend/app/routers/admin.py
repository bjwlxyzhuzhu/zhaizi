import io
import re
from datetime import datetime

from fastapi import APIRouter, Depends, File, Form, HTTPException, Request, UploadFile
from fastapi.responses import FileResponse, StreamingResponse
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from pydantic import BaseModel
from sqlalchemy import func, or_, select
from sqlalchemy.orm import Session

from .. import points
from ..config import EXPORT_DIR
from ..db import get_db
from ..deps import ROLE_LABEL, audit, client_ip, require, require_super
from ..models import (AuditLog, ImportJob, KPITarget, PointLedger, PointRule, SystemSetting, TaskTemplate,
                      TrustSample, User)
from ..qa_engine import DEFAULT_SENSITIVE
from ..security import encrypt_secret, decrypt_secret, hash_password, random_password

router = APIRouter(prefix="/admin", tags=["管理后台"])
admin_only = require("admin")
IDENTITIES = ["教师", "研究生", "本科生", "文博人员", "中小学教师", "其他"]
ROLE_FROM_CN = {"用户": "user", "研究员": "researcher", "管理员": "admin"}
ACCOUNT_RE = re.compile(r"^(1\d{10}|[A-Za-z0-9]{4,20})$")
EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")


def _u(db: Session, u: User) -> dict:
    return {"id": u.id, "account": u.account, "name": u.name, "org": u.org, "identity_type": u.identity_type,
            "role": u.role, "role_label": ROLE_LABEL[u.role], "is_super": u.is_super, "status": u.status,
            "phone": u.phone, "email": u.email, "remark": u.remark, "points": points.balance(db, u.id),
            "locked": bool(u.locked_until and u.locked_until > datetime.now()),
            "must_change_pwd": u.must_change_pwd,
            "last_login_at": u.last_login_at.isoformat(timespec="minutes") if u.last_login_at else "",
            "created_at": u.created_at.isoformat(timespec="minutes")}


# ---------------- 用户管理 ----------------
@router.get("/users")
def list_users(q: str = "", org: str = "", role: str = "", status: str = "", page: int = 1, size: int = 20,
               me: User = Depends(admin_only), db: Session = Depends(get_db)):
    stmt = select(User)
    if q:
        stmt = stmt.where(or_(User.name.contains(q), User.account.contains(q)))
    if org:
        stmt = stmt.where(User.org.contains(org))
    if role:
        stmt = stmt.where(User.role == role)
    if status:
        stmt = stmt.where(User.status == status)
    total = db.scalar(select(func.count()).select_from(stmt.subquery()))
    rows = db.scalars(stmt.order_by(User.id.desc()).offset((page - 1) * size).limit(size)).all()
    orgs = [o for o in db.scalars(select(User.org).distinct()) if o]
    return {"total": total, "items": [_u(db, u) for u in rows], "orgs": orgs}


class UserIn(BaseModel):
    account: str
    name: str
    org: str = ""
    identity_type: str = "其他"
    role: str = "user"
    status: str = "active"
    phone: str = ""
    email: str = ""
    remark: str = ""
    initial_points: int | None = None


@router.post("/users")
def create_user(body: UserIn, request: Request, me: User = Depends(admin_only), db: Session = Depends(get_db)):
    if not ACCOUNT_RE.match(body.account):
        raise HTTPException(400, "账号须为学工号（4-20位字母数字）或 11 位手机号")
    if db.scalar(select(User).where(User.account == body.account)):
        raise HTTPException(400, "账号已存在")
    if body.role == "admin" and not me.is_super:
        raise HTTPException(403, "设为管理员需超级管理员操作")
    pwd = random_password()
    u = User(account=body.account, name=body.name, org=body.org, identity_type=body.identity_type, role=body.role,
             status=body.status, phone=body.phone, email=body.email, remark=body.remark,
             password_hash=hash_password(pwd), must_change_pwd=True)
    db.add(u)
    db.flush()
    init = body.initial_points if body.initial_points is not None else (points.rule(db, "new_user_default") if body.role == "user" else 0)
    if init:
        points.add_entry(db, u.id, "GRANT", init, reason="新用户初始积分", operator_id=me.id)
    audit(db, me, "新增用户", body.account, {"role": body.role}, client_ip(request))
    db.commit()
    return {"user": _u(db, u), "initial_password": pwd}


class UserPatch(BaseModel):
    name: str | None = None
    org: str | None = None
    identity_type: str | None = None
    role: str | None = None
    status: str | None = None
    phone: str | None = None
    email: str | None = None
    remark: str | None = None


@router.patch("/users/{uid}")
def edit_user(uid: int, body: UserPatch, request: Request, me: User = Depends(admin_only), db: Session = Depends(get_db)):
    u = db.get(User, uid) or _404()
    data = body.model_dump(exclude_none=True)
    if "role" in data and data["role"] != u.role:
        if (data["role"] == "admin" or u.role == "admin") and not me.is_super:
            raise HTTPException(403, "设置或撤销管理员需超级管理员操作")
        if u.is_super:
            raise HTTPException(400, "不能变更超级管理员角色")
        audit(db, me, "角色变更", u.account, {"from": u.role, "to": data["role"]}, client_ip(request))
        u.token_version += 1  # 下次请求生效
    if "status" in data and data["status"] != u.status:
        if u.is_super:
            raise HTTPException(400, "不能停用超级管理员")
        audit(db, me, "启用" if data["status"] == "active" else "停用", u.account, {}, client_ip(request))
        u.token_version += 1  # AUTH-6 停用立即失效
    for k, v in data.items():
        setattr(u, k, v)
    audit(db, me, "编辑用户", u.account, {k: v for k, v in data.items() if k not in ("role", "status")}, client_ip(request))
    db.commit()
    return _u(db, u)


@router.post("/users/{uid}/reset-password")
def reset_password(uid: int, request: Request, me: User = Depends(admin_only), db: Session = Depends(get_db)):
    u = db.get(User, uid) or _404()
    if u.is_super and not me.is_super:
        raise HTTPException(403, "无权重置超级管理员密码")
    pwd = random_password()
    u.password_hash = hash_password(pwd)
    u.must_change_pwd = True
    u.token_version += 1
    u.locked_until = None
    u.failed_attempts = 0
    audit(db, me, "重置密码", u.account, {}, client_ip(request))
    db.commit()
    return {"initial_password": pwd}


@router.post("/users/{uid}/unlock")
def unlock(uid: int, request: Request, me: User = Depends(admin_only), db: Session = Depends(get_db)):
    u = db.get(User, uid) or _404()
    u.locked_until = None
    u.failed_attempts = 0
    audit(db, me, "解锁账号", u.account, {}, client_ip(request))
    db.commit()
    return {"ok": True}


class BatchIn(BaseModel):
    ids: list[int]
    action: str  # role / enable / disable / grant
    role: str | None = None
    amount: int | None = None
    reason: str = ""


@router.post("/users/batch")
def batch(body: BatchIn, request: Request, me: User = Depends(admin_only), db: Session = Depends(get_db)):
    users = db.scalars(select(User).where(User.id.in_(body.ids))).all()
    n = 0
    for u in users:
        if u.is_super:
            continue
        if body.action == "role" and body.role:
            if (body.role == "admin" or u.role == "admin") and not me.is_super:
                continue
            u.role, u.token_version = body.role, u.token_version + 1
        elif body.action in ("enable", "disable"):
            u.status = "active" if body.action == "enable" else "disabled"
            u.token_version += 1
        elif body.action == "grant":
            if not body.amount or body.amount <= 0 or not body.reason:
                raise HTTPException(400, "批量发放须填写正整数积分与原因")
            points.add_entry(db, u.id, "GRANT", body.amount, reason=body.reason, operator_id=me.id)
        n += 1
    audit(db, me, "批量操作", f"{n} 人", body.model_dump(), client_ip(request))
    db.commit()
    return {"affected": n}


def _404():
    raise HTTPException(404, "不存在")


# ---------------- Excel 导入（ADMIN-5~10，附录A） ----------------
HEADERS = ["账号", "姓名", "单位", "身份类别", "角色", "初始积分", "手机号", "邮箱", "备注"]


@router.get("/import/template")
def template(me: User = Depends(admin_only)):
    wb = Workbook()
    ws = wb.active
    ws.title = "用户导入"
    ws.append(HEADERS)
    ws.append(["2026001", "张三", "人文学院", "研究生", "用户", 100, "", "", "示例行，导入前请删除"])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="3F8F7A")
    for col, w in zip("ABCDEFGHI", [16, 10, 24, 12, 10, 10, 14, 22, 24]):
        ws.column_dimensions[col].width = w
    dv1 = DataValidation(type="list", formula1='"' + ",".join(IDENTITIES) + '"', allow_blank=False)
    dv2 = DataValidation(type="list", formula1='"用户,研究员"', allow_blank=False)
    ws.add_data_validation(dv1)
    ws.add_data_validation(dv2)
    dv1.add("D2:D2001")
    dv2.add("E2:E2001")
    help_ws = wb.create_sheet("填写说明")
    rows = [("列", "字段", "必填", "规则"),
            ("A", "账号", "是", "学工号或11位手机号，文件内不得重复"),
            ("B", "姓名", "是", "≤20字"),
            ("C", "单位", "是", "如“人文学院”“某博物馆”"),
            ("D", "身份类别", "是", "下拉：" + "/".join(IDENTITIES)),
            ("E", "角色", "是", "下拉：用户/研究员（管理员不可通过导入设置）"),
            ("F", "初始积分", "否", "非负整数；留空按默认值；研究员可留空"),
            ("G", "手机号", "否", "11位"),
            ("H", "邮箱", "否", "邮箱格式"),
            ("I", "备注", "否", "≤100字"),
            ("", "", "", "单次最多 2000 行。已存在账号可选择跳过或更新信息（不覆盖密码）。")]
    for r in rows:
        help_ws.append(r)
    help_ws.column_dimensions["D"].width = 60
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": "attachment; filename*=UTF-8''%E7%94%A8%E6%88%B7%E5%AF%BC%E5%85%A5%E6%A8%A1%E6%9D%BF.xlsx"})


def _cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


@router.post("/import/preview")
async def import_preview(file: UploadFile = File(...), mode: str = Form("skip"), me: User = Depends(admin_only),
                         db: Session = Depends(get_db)):
    try:
        wb = load_workbook(io.BytesIO(await file.read()), read_only=True, data_only=True)
    except Exception:
        raise HTTPException(400, "无法读取文件，请使用模板 .xlsx 格式")
    ws = wb.worksheets[0]
    raw_rows = []
    for idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        vals = [_cell(v) for v in (list(row) + [None] * 9)[:9]]
        if not any(vals):
            continue
        raw_rows.append((idx, vals))
    if len(raw_rows) > 2000:
        raise HTTPException(400, f"单次最多导入 2000 行，当前 {len(raw_rows)} 行")
    existing = {a for (a,) in db.execute(select(User.account))}
    seen: dict[str, int] = {}
    default_pts = points.rule(db, "new_user_default")
    rows_out = []
    for line, v in raw_rows:
        account, name, org, ident, role_cn, pts, phone, email, remark = v
        errs = []
        if not account:
            errs.append("账号必填")
        elif not ACCOUNT_RE.match(account):
            errs.append("账号格式应为学工号（4-20位字母数字）或11位手机号")
        if account in seen:
            errs.append(f"与第 {seen[account]} 行账号重复")
        if account:
            seen.setdefault(account, line)
        if not name:
            errs.append("姓名必填")
        elif len(name) > 20:
            errs.append("姓名超过20字")
        if not org:
            errs.append("单位必填")
        if not ident:
            errs.append("身份类别必填")
        elif ident not in IDENTITIES:
            errs.append("身份类别取值无效")
        if not role_cn:
            errs.append("角色必填")
        elif role_cn == "管理员":
            errs.append("管理员不可通过导入设置")
        elif role_cn not in ("用户", "研究员"):
            errs.append("角色只能为“用户”或“研究员”")
        if pts and not re.fullmatch(r"\d+", pts):
            errs.append("初始积分须为非负整数")
        if phone and not re.fullmatch(r"1\d{10}", phone):
            errs.append("手机号须为11位")
        if email and not EMAIL_RE.match(email):
            errs.append("邮箱格式不正确")
        if len(remark) > 100:
            errs.append("备注超过100字")
        exists = account in existing
        action = "error" if errs else ("update" if exists and mode == "update" else ("skip" if exists else "create"))
        init = int(pts) if pts else (default_pts if role_cn == "用户" else 0)
        rows_out.append({"line": line, "raw": v, "errors": errs, "action": action,
                         "data": {"account": account, "name": name, "org": org, "identity_type": ident,
                                  "role": ROLE_FROM_CN.get(role_cn, "user"), "points": init,
                                  "points_given": bool(pts), "phone": phone, "email": email, "remark": remark}})
    job = ImportJob(operator_id=me.id, filename=file.filename or "导入.xlsx", mode=mode, total=len(rows_out),
                    created=sum(r["action"] == "create" for r in rows_out),
                    updated=sum(r["action"] == "update" for r in rows_out),
                    failed=sum(r["action"] == "error" for r in rows_out), payload={"rows": rows_out})
    db.add(job)
    db.commit()
    return {"job_id": job.id, "total": job.total, "create": job.created, "update": job.updated,
            "skip": sum(r["action"] == "skip" for r in rows_out), "error": job.failed, "rows": rows_out}


@router.post("/import/{job_id}/commit")
def import_commit(job_id: int, request: Request, me: User = Depends(admin_only), db: Session = Depends(get_db)):
    job = db.get(ImportJob, job_id) or _404()
    if job.status != "preview":
        raise HTTPException(400, "该导入任务已提交")
    rows = job.payload.get("rows", [])
    created, updated, errors = [], 0, []
    hashed_cache = {}
    for r in rows:
        d = r["data"]
        if r["action"] == "error":
            errors.append(r)
            continue
        if r["action"] == "create":
            pwd = random_password()
            hashed_cache[d["account"]] = pwd
            u = User(account=d["account"], name=d["name"], org=d["org"], identity_type=d["identity_type"],
                     role=d["role"], phone=d["phone"], email=d["email"], remark=d["remark"],
                     password_hash=hash_password(pwd), must_change_pwd=True)
            db.add(u)
            db.flush()
            if d["points"]:
                points.add_entry(db, u.id, "IMPORT", d["points"], reason=f"导入发放（{job.filename}）",
                                 ref_type="import", ref_id=str(job.id), operator_id=me.id)
            created.append((d, pwd))
        elif r["action"] == "update":
            u = db.scalar(select(User).where(User.account == d["account"]))
            if not u:
                continue
            for k in ("name", "org", "identity_type", "phone", "email", "remark"):
                if d[k]:
                    setattr(u, k, d[k])
            if u.role != "admin" and d["role"] != u.role:
                u.role, u.token_version = d["role"], u.token_version + 1
            if d["points_given"] and d["points"]:
                # ADMIN-8：更新时积分记一笔“导入发放”流水，而不是直接改余额
                points.add_entry(db, u.id, "IMPORT", d["points"], reason=f"导入发放（{job.filename}）",
                                 ref_type="import", ref_id=str(job.id), operator_id=me.id)
            updated += 1
    # 结果文件：成功清单（含初始密码，仅可下载一次）与错误清单
    wb = Workbook()
    ws = wb.active
    ws.title = "成功清单"
    ws.append(["账号", "姓名", "单位", "角色", "初始密码", "初始积分"])
    for d, pwd in created:
        ws.append([d["account"], d["name"], d["org"], ROLE_LABEL[d["role"]], pwd, d["points"]])
    ok_name = f"import_{job.id}_success.xlsx"
    wb.save(EXPORT_DIR / ok_name)
    ewb = Workbook()
    ews = ewb.active
    ews.title = "错误清单"
    ews.append(["原行号"] + HEADERS + ["错误原因"])
    for r in errors:
        ews.append([r["line"]] + r["raw"] + ["；".join(r["errors"])])
    err_name = f"import_{job.id}_errors.xlsx"
    ewb.save(EXPORT_DIR / err_name)
    job.status, job.created, job.updated, job.failed = "done", len(created), updated, len(errors)
    job.result_file, job.error_file = ok_name, err_name
    job.payload = {"rows": [{**r, "data": {**r["data"]}} for r in rows]}  # 不保存明文密码
    audit(db, me, "Excel导入", job.filename, {"created": len(created), "updated": updated, "failed": len(errors)},
          client_ip(request))
    db.commit()
    return {"created": len(created), "updated": updated, "failed": len(errors), "job_id": job.id,
            "accounts": [{"account": d["account"], "name": d["name"], "password": p} for d, p in created]}


@router.get("/import/{job_id}/result")
def import_result(job_id: int, kind: str = "success", me: User = Depends(admin_only), db: Session = Depends(get_db)):
    job = db.get(ImportJob, job_id) or _404()
    if kind == "success":
        if job.result_downloaded:
            raise HTTPException(410, "初始密码清单仅可下载一次，已下载过")
        path = EXPORT_DIR / job.result_file
        job.result_downloaded = True
        db.commit()
        return FileResponse(path, filename=f"导入成功清单_{job.id}.xlsx")
    return FileResponse(EXPORT_DIR / job.error_file, filename=f"导入错误清单_{job.id}.xlsx")


@router.get("/import/jobs")
def import_jobs(me: User = Depends(admin_only), db: Session = Depends(get_db)):
    rows = db.scalars(select(ImportJob).where(ImportJob.status == "done").order_by(ImportJob.id.desc()).limit(100)).all()
    names = {u.id: u.name for u in db.scalars(select(User).where(User.id.in_({j.operator_id for j in rows})))}
    return [{"id": j.id, "operator": names.get(j.operator_id, ""), "filename": j.filename, "total": j.total,
             "created": j.created, "updated": j.updated, "failed": j.failed, "result_downloaded": j.result_downloaded,
             "created_at": j.created_at.isoformat(timespec="minutes")} for j in rows]


# ---------------- 积分管理 ----------------
@router.get("/points/ledger")
def ledger(q: str = "", type: str = "", start: str = "", end: str = "", page: int = 1, size: int = 20,
           export: bool = False, me: User = Depends(admin_only), db: Session = Depends(get_db)):
    stmt = select(PointLedger, User).join(User, User.id == PointLedger.user_id)
    if q:
        stmt = stmt.where(or_(User.name.contains(q), User.account.contains(q)))
    if type:
        stmt = stmt.where(PointLedger.type == type)
    if start:
        stmt = stmt.where(PointLedger.created_at >= datetime.fromisoformat(start))
    if end:
        stmt = stmt.where(PointLedger.created_at <= datetime.fromisoformat(end + "T23:59:59"))
    stmt = stmt.order_by(PointLedger.id.desc())
    if export:
        wb = Workbook()
        ws = wb.active
        ws.append(["流水号", "账号", "姓名", "类型", "变动", "变动后余额", "原因", "时间"])
        for e, u in db.execute(stmt.limit(20000)):
            ws.append([e.id, u.account, u.name, e.type, e.amount, e.balance_after, e.reason,
                       e.created_at.strftime("%Y-%m-%d %H:%M:%S")])
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                                 headers={"Content-Disposition": "attachment; filename=ledger.xlsx"})
    total = db.scalar(select(func.count()).select_from(stmt.subquery()))
    rows = db.execute(stmt.offset((page - 1) * size).limit(size)).all()
    return {"total": total, "items": [{"id": e.id, "account": u.account, "name": u.name, "type": e.type,
                                       "amount": e.amount, "balance_after": e.balance_after, "reason": e.reason,
                                       "created_at": e.created_at.isoformat(timespec="seconds")} for e, u in rows]}


class AdjustIn(BaseModel):
    user_id: int
    amount: int
    reason: str


@router.post("/points/adjust")
def adjust(body: AdjustIn, request: Request, me: User = Depends(admin_only), db: Session = Depends(get_db)):
    if not body.reason.strip():
        raise HTTPException(400, "必须填写调整原因")
    if body.amount == 0:
        raise HTTPException(400, "调整数值不能为 0")
    u = db.get(User, body.user_id) or _404()
    with points.user_lock(u.id):
        e = points.add_entry(db, u.id, "ADJUST", body.amount, reason=body.reason, operator_id=me.id)
        audit(db, me, "积分调整", u.account, {"amount": body.amount, "reason": body.reason}, client_ip(request))
        db.commit()
    return {"balance": e.balance_after}


class GrantIn(BaseModel):
    user_ids: list[int] = []
    org: str = ""
    role: str = ""
    amount: int
    reason: str


@router.post("/points/grant")
def grant(body: GrantIn, request: Request, me: User = Depends(admin_only), db: Session = Depends(get_db)):
    if body.amount <= 0 or not body.reason.strip():
        raise HTTPException(400, "须填写正整数积分与原因")
    stmt = select(User).where(User.status == "active")
    if body.user_ids:
        stmt = stmt.where(User.id.in_(body.user_ids))
    else:
        if body.org:
            stmt = stmt.where(User.org == body.org)
        if body.role:
            stmt = stmt.where(User.role == body.role)
    users = db.scalars(stmt).all()
    for u in users:
        points.add_entry(db, u.id, "GRANT", body.amount, reason=body.reason, operator_id=me.id)
    audit(db, me, "批量发放积分", f"{len(users)} 人", {"amount": body.amount, "reason": body.reason}, client_ip(request))
    db.commit()
    return {"affected": len(users)}


@router.get("/point-rules")
def get_rules(me: User = Depends(admin_only), db: Session = Depends(get_db)):
    out = []
    for key, label, default in points.DEFAULT_RULES:
        r = db.get(PointRule, key)
        out.append({"key": key, "label": label, "value": r.value if r else default,
                    "updated_at": r.updated_at.isoformat(timespec="minutes") if r else ""})
    return out


@router.put("/point-rules")
def put_rules(body: dict[str, int], request: Request, me: User = Depends(admin_only), db: Session = Depends(get_db)):
    valid = {k for k, _, _ in points.DEFAULT_RULES}
    changes = {}
    for k, v in body.items():
        if k not in valid or v < 0:
            raise HTTPException(400, f"规则 {k} 无效")
        r = db.get(PointRule, k)
        label = next(l for kk, l, _ in points.DEFAULT_RULES if kk == k)
        if not r:
            r = PointRule(key=k, label=label, value=v)
            db.add(r)
        if r.value != v or not r.updated_by:
            changes[k] = {"from": r.value, "to": v}
        r.value, r.updated_by, r.updated_at = v, me.id, datetime.now()
    audit(db, me, "积分规则修改", "point_rules", changes, client_ip(request))
    db.commit()
    return {"ok": True}


# ---------------- 任务模板 ----------------
class TemplateIn(BaseModel):
    code: str
    name: str
    description: str = ""
    reward: int = 0
    need_review: bool = True
    daily_limit: int = 20
    enabled: bool = True


@router.get("/task-templates")
def templates(me: User = Depends(require("admin", "researcher")), db: Session = Depends(get_db)):
    return [{"id": t.id, "code": t.code, "name": t.name, "description": t.description, "reward": t.reward,
             "need_review": t.need_review, "daily_limit": t.daily_limit, "enabled": t.enabled}
            for t in db.scalars(select(TaskTemplate))]


@router.post("/task-templates")
def upsert_template(body: TemplateIn, request: Request, me: User = Depends(admin_only), db: Session = Depends(get_db)):
    t = db.scalar(select(TaskTemplate).where(TaskTemplate.code == body.code))
    if not t:
        t = TaskTemplate(code=body.code)
        db.add(t)
    for k, v in body.model_dump().items():
        setattr(t, k, v)
    audit(db, me, "任务模板修改", body.name, body.model_dump(), client_ip(request))
    db.commit()
    return {"ok": True}


# ---------------- 审计 ----------------
@router.get("/audit-logs")
def audit_logs(action: str = "", q: str = "", page: int = 1, size: int = 30, me: User = Depends(admin_only),
               db: Session = Depends(get_db)):
    stmt = select(AuditLog)
    if action:
        stmt = stmt.where(AuditLog.action == action)
    if q:
        stmt = stmt.where(or_(AuditLog.actor_name.contains(q), AuditLog.target.contains(q)))
    total = db.scalar(select(func.count()).select_from(stmt.subquery()))
    rows = db.scalars(stmt.order_by(AuditLog.id.desc()).offset((page - 1) * size).limit(size)).all()
    actions = [a for (a,) in db.execute(select(AuditLog.action).distinct())]
    return {"total": total, "actions": actions,
            "items": [{"id": a.id, "actor": a.actor_name, "action": a.action, "target": a.target, "detail": a.detail,
                       "ip": a.ip, "created_at": a.created_at.isoformat(timespec="seconds")} for a in rows]}


# ---------------- 系统设置（仅超级管理员，ADMIN-17） ----------------
def _setting(db, key, default=""):
    s = db.get(SystemSetting, key)
    return s.value if s else default


def _set(db, key, value):
    s = db.get(SystemSetting, key)
    if not s:
        s = SystemSetting(key=key)
        db.add(s)
    s.value, s.updated_at = value, datetime.now()


@router.get("/settings")
def get_settings(me: User = Depends(require_super), db: Session = Depends(get_db)):
    from .. import config
    key = decrypt_secret(_setting(db, "model_api_key_enc")) or config.MODEL_API_KEY
    return {"model_api_key_tail": ("****" + key[-4:]) if key else "", "model_base_url": _setting(db, "model_base_url", config.MODEL_BASE_URL),
            "model_name": _setting(db, "model_name", config.MODEL_NAME),
            "model_timeout": int(_setting(db, "model_timeout", str(config.MODEL_TIMEOUT))),
            "price_in": float(_setting(db, "price_in", "0.0024")), "price_out": float(_setting(db, "price_out", "0.0096")),
            "allow_l2_cloud": False, "allow_l2_cloud_locked": True,
            "sensitive_words": _setting(db, "sensitive_words", ",".join(DEFAULT_SENSITIVE))}


class SettingsIn(BaseModel):
    model_api_key: str | None = None
    model_base_url: str | None = None
    model_name: str | None = None
    model_timeout: int | None = None
    price_in: float | None = None
    price_out: float | None = None
    sensitive_words: str | None = None


@router.put("/settings")
def put_settings(body: SettingsIn, request: Request, me: User = Depends(require_super), db: Session = Depends(get_db)):
    changed = []
    if body.model_api_key is not None:
        key = body.model_api_key.strip()
        _set(db, "model_api_key_enc", encrypt_secret(key) if key else "")
        changed.append("model_api_key")
    for k in ("model_base_url", "model_name", "model_timeout", "price_in", "price_out", "sensitive_words"):
        v = getattr(body, k)
        if v is not None:
            _set(db, k, str(v))
            changed.append(k)
    audit(db, me, "系统设置修改", "system_settings", {"fields": changed}, client_ip(request))
    db.commit()
    return {"ok": True}


# ---------------- 考核指标目标值（DASH-7）与抽检记录 ----------------
@router.get("/kpi-targets")
def kpi_targets(me: User = Depends(require("admin", "researcher")), db: Session = Depends(get_db)):
    return [{"id": k.id, "indicator": k.indicator, "metric": k.metric, "year": k.year, "target": k.target, "unit": k.unit}
            for k in db.scalars(select(KPITarget).order_by(KPITarget.year, KPITarget.id))]


@router.put("/kpi-targets/{kid}")
def put_kpi(kid: int, body: dict, request: Request, me: User = Depends(admin_only), db: Session = Depends(get_db)):
    k = db.get(KPITarget, kid) or _404()
    k.target = float(body["target"])
    audit(db, me, "考核目标修改", f"{k.indicator}({k.year})", {"target": k.target}, client_ip(request))
    db.commit()
    return {"ok": True}


class SampleIn(BaseModel):
    sampled: int
    accurate: int
    claims: int
    unsourced_claims: int
    note: str = ""


@router.post("/trust-samples")
def add_sample(body: SampleIn, request: Request, me: User = Depends(require("admin", "researcher")),
               db: Session = Depends(get_db)):
    if body.accurate > body.sampled or body.unsourced_claims > body.claims:
        raise HTTPException(400, "数值不合理")
    db.add(TrustSample(**body.model_dump(), researcher_id=me.id))
    audit(db, me, "录入抽检记录", "trust_samples", body.model_dump(), client_ip(request))
    db.commit()
    return {"ok": True}
