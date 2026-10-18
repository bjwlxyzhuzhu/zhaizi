"""研究员工作台：数据录入、双人审核、版本、裁断、任务审核（PRD §5.5）。"""
import io
import uuid
from datetime import datetime

from fastapi import APIRouter, Depends, File, HTTPException, Request, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.datavalidation import DataValidation
from pydantic import BaseModel
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from .. import points, retrieval
from ..config import UPLOAD_DIR
from ..db import get_db
from ..deps import audit, client_ip, require
from ..kg_builder import link_record
from ..models import (Adjudication, DialectEntry, Glyph, Inscription, Notification, RecordVersion, TaskBatch,
                      TaskSubmission, TaskTemplate, TextPassage, User, Vessel)

router = APIRouter(prefix="/res", tags=["研究员"])
researcher = require("researcher")
res_or_admin = require("researcher", "admin")

TABLES = {
    "vessels": (Vessel, ["name", "category", "period", "find_place", "district", "find_date", "collection",
                         "catalog_refs", "description", "image", "image_credit", "level", "evidence_grade", "source"]),
    "inscriptions": (Inscription, ["vessel_id", "transcription", "interpretation", "char_count", "source",
                                   "evidence_grade", "level"]),
    "glyphs": (Glyph, ["character", "script", "vessel_id", "image_path", "note", "source", "level"]),
    "dialect_entries": (DialectEntry, ["district", "speaker_code", "word", "ipa", "meaning", "audio_path", "source", "level"]),
    "texts": (TextPassage, ["title", "chapter", "content", "source", "related_vessel_id", "level"]),
}
TABLE_LABEL = {"vessels": "器物", "inscriptions": "铭文", "glyphs": "字形", "dialect_entries": "方言条目", "texts": "传世文献"}
TITLE_FIELD = {"vessels": "name", "inscriptions": "transcription", "glyphs": "character", "dialect_entries": "word", "texts": "title"}


def _model(table: str):
    if table not in TABLES:
        raise HTTPException(400, "未知数据类型")
    return TABLES[table]


def snapshot(obj, fields) -> dict:
    return {f: getattr(obj, f) for f in fields} | {"status": obj.status}


def save_version(db, table, obj, fields, editor, action):
    n = db.scalar(select(func.count()).select_from(RecordVersion).where(
        RecordVersion.table_name == table, RecordVersion.record_id == obj.id)) or 0
    db.add(RecordVersion(table_name=table, record_id=obj.id, version=n + 1, data=snapshot(obj, fields),
                         action=action, editor_id=editor.id if editor else None))


def _row(table, obj, names) -> dict:
    _, fields = TABLES[table]
    title = str(getattr(obj, TITLE_FIELD[table]) or "")
    return {"id": obj.id, "table": table, "table_label": TABLE_LABEL[table], "title": title[:40],
            "status": obj.status, "level": obj.level, "source": getattr(obj, "source", ""),
            "created_by": names.get(obj.created_by, "种子数据"), "created_by_id": obj.created_by,
            "reviewed_by": names.get(obj.reviewed_by, ""), "updated_at": obj.updated_at.isoformat(timespec="minutes"),
            "data": {f: getattr(obj, f) for f in fields}}


@router.get("/records")
def records(table: str = "vessels", status: str = "", mine: bool = False, q: str = "",
            u: User = Depends(res_or_admin), db: Session = Depends(get_db)):
    M, _ = _model(table)
    stmt = select(M)
    if status:
        stmt = stmt.where(M.status == status)
    if mine:
        stmt = stmt.where(M.created_by == u.id)
    if q:
        stmt = stmt.where(getattr(M, TITLE_FIELD[table]).contains(q))
    rows = db.scalars(stmt.order_by(M.updated_at.desc()).limit(300)).all()
    names = {x.id: x.name for x in db.scalars(select(User).where(User.role.in_(("researcher", "admin"))))}
    counts = {s: db.scalar(select(func.count()).select_from(M).where(M.status == s)) for s in ("pending", "approved", "rejected")}
    return {"items": [_row(table, r, names) for r in rows], "counts": counts}


class RecordIn(BaseModel):
    table: str
    data: dict


def _clean(table, data: dict) -> dict:
    _, fields = TABLES[table]
    out = {k: v for k, v in data.items() if k in fields}
    if not str(out.get("source", "")).strip():
        raise HTTPException(400, "每条记录必须填写来源")  # RES-1
    if out.get("level") not in ("L1", "L2", "L3"):
        raise HTTPException(400, "必须选择数据分级 L1/L2/L3")
    for k in ("vessel_id", "related_vessel_id", "char_count"):
        if k in out:
            out[k] = int(out[k]) if str(out[k] or "").strip() else None
    if table == "inscriptions" and out.get("transcription") and not out.get("char_count"):
        out["char_count"] = sum(1 for c in out["transcription"] if "㐀" <= c <= "鿿")
    return out


@router.post("/records")
def create_record(body: RecordIn, request: Request, u: User = Depends(researcher), db: Session = Depends(get_db)):
    M, fields = _model(body.table)
    obj = M(**_clean(body.table, body.data), status="pending", created_by=u.id)
    db.add(obj)
    db.flush()
    save_version(db, body.table, obj, fields, u, "create")
    retrieval.SYNC[body.table](db, obj)  # 以 pending 状态登记检索块，审核通过前不参与检索
    audit(db, u, "数据录入", f"{TABLE_LABEL[body.table]}#{obj.id}", {}, client_ip(request))
    db.commit()
    return {"id": obj.id}


@router.put("/records/{table}/{rid}")
def update_record(table: str, rid: int, body: RecordIn, request: Request, u: User = Depends(researcher),
                  db: Session = Depends(get_db)):
    M, fields = _model(table)
    obj = db.get(M, rid) or _404()
    for k, v in _clean(table, body.data).items():
        setattr(obj, k, v)
    obj.status, obj.updated_at = "pending", datetime.now()  # 修改后重新进入待审
    obj.created_by = obj.created_by or u.id
    save_version(db, table, obj, fields, u, "edit")
    retrieval.SYNC[table](db, obj)
    audit(db, u, "数据修改", f"{TABLE_LABEL[table]}#{rid}", {}, client_ip(request))
    db.commit()
    return {"ok": True}


class ReviewIn(BaseModel):
    approve: bool
    reason: str = ""


@router.post("/records/{table}/{rid}/review")
def review_record(table: str, rid: int, body: ReviewIn, request: Request, u: User = Depends(researcher),
                  db: Session = Depends(get_db)):
    M, fields = _model(table)
    obj = db.get(M, rid) or _404()
    if obj.created_by == u.id:
        raise HTTPException(403, "双人原则：不能审核本人录入的数据")  # RES-2 / T7
    if obj.status != "pending":
        raise HTTPException(400, "该记录不在待审状态")
    if not body.approve and not body.reason.strip():
        raise HTTPException(400, "驳回须填写原因")
    obj.status = "approved" if body.approve else "rejected"
    obj.reviewed_by, obj.updated_at = u.id, datetime.now()
    save_version(db, table, obj, fields, u, "approve" if body.approve else "reject")
    retrieval.SYNC[table](db, obj)
    if body.approve:
        link_record(db, table, obj)  # 自动生成/更新图谱节点
    if obj.created_by:
        db.add(Notification(user_id=obj.created_by, kind="review",
                            title=f"{TABLE_LABEL[table]}“{str(getattr(obj, TITLE_FIELD[table]))[:12]}”" + ("已审核通过" if body.approve else "被驳回"),
                            body=body.reason))
    audit(db, u, "数据审核", f"{TABLE_LABEL[table]}#{rid}", {"approve": body.approve, "reason": body.reason}, client_ip(request))
    db.commit()
    return {"ok": True}


@router.get("/records/{table}/{rid}/versions")
def versions(table: str, rid: int, u: User = Depends(res_or_admin), db: Session = Depends(get_db)):
    rows = db.scalars(select(RecordVersion).where(RecordVersion.table_name == table, RecordVersion.record_id == rid)
                      .order_by(RecordVersion.version.desc())).all()
    names = {x.id: x.name for x in db.scalars(select(User))}
    return [{"version": r.version, "action": r.action, "data": r.data, "editor": names.get(r.editor_id, "种子数据"),
             "created_at": r.created_at.isoformat(timespec="minutes")} for r in rows]


@router.post("/records/{table}/{rid}/rollback/{version}")
def rollback(table: str, rid: int, version: int, request: Request, u: User = Depends(researcher),
             db: Session = Depends(get_db)):
    M, fields = _model(table)
    obj = db.get(M, rid) or _404()
    v = db.scalar(select(RecordVersion).where(RecordVersion.table_name == table, RecordVersion.record_id == rid,
                                              RecordVersion.version == version)) or _404()
    for f in fields:
        if f in v.data:
            setattr(obj, f, v.data[f])
    obj.status, obj.updated_at = "pending", datetime.now()
    save_version(db, table, obj, fields, u, f"rollback→v{version}")
    retrieval.SYNC[table](db, obj)
    audit(db, u, "版本回退", f"{TABLE_LABEL[table]}#{rid}", {"to": version}, client_ip(request))
    db.commit()
    return {"ok": True}


@router.post("/upload")
async def upload(file: UploadFile = File(...), u: User = Depends(researcher)):
    data = await file.read()
    if len(data) > 20 * 1024 * 1024:
        raise HTTPException(400, "文件过大")
    ext = (file.filename or "bin").rsplit(".", 1)[-1].lower()
    if ext not in ("png", "jpg", "jpeg", "svg", "webp", "mp3", "wav", "m4a"):
        raise HTTPException(400, "不支持的文件类型")
    rel = f"res/{uuid.uuid4().hex}.{ext}"
    (UPLOAD_DIR / "res").mkdir(exist_ok=True)
    (UPLOAD_DIR / rel).write_bytes(data)
    return {"path": rel, "url": f"/api/files/{rel}"}


# ---------------- 器铭 / 方言 Excel 批量录入 ----------------
IMPORT_COLS = {
    "vessels": ["器名", "类别", "时代", "出土地", "所属区县", "出土时间", "现藏地", "著录号", "简介", "证据等级", "数据分级", "来源", "铭文释文"],
    "dialect_entries": ["区县", "发音人编号", "词条", "国际音标", "词义", "来源", "数据分级"],
}


@router.get("/import/template")
def res_template(kind: str = "vessels", u: User = Depends(researcher)):
    wb = Workbook()
    ws = wb.active
    ws.title = "录入"
    ws.append(IMPORT_COLS[kind])
    dv = DataValidation(type="list", formula1='"L1,L2,L3"')
    ws.add_data_validation(dv)
    col = "K" if kind == "vessels" else "G"
    dv.add(f"{col}2:{col}2001")
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f"attachment; filename={kind}_template.xlsx"})


@router.post("/import")
async def res_import(kind: str, request: Request, file: UploadFile = File(...), u: User = Depends(researcher),
                     db: Session = Depends(get_db)):
    wb = load_workbook(io.BytesIO(await file.read()), read_only=True, data_only=True)
    ok, errors = 0, []
    for line, row in enumerate(wb.worksheets[0].iter_rows(min_row=2, values_only=True), start=2):
        vals = [("" if v is None else str(v).strip()) for v in (list(row) + [None] * 13)[:13]]
        if not any(vals):
            continue
        try:
            if kind == "vessels":
                d = dict(zip(TABLES["vessels"][1][:9], vals[:9]))
                d.update(evidence_grade=vals[9] or "B", level=vals[10], source=vals[11])
                if not d["name"]:
                    raise HTTPException(400, "器名必填")
                v = Vessel(**_clean("vessels", d), status="pending", created_by=u.id)
                db.add(v)
                db.flush()
                save_version(db, "vessels", v, TABLES["vessels"][1], u, "import")
                retrieval.sync_vessel(db, v)
                if vals[12]:
                    i = Inscription(**_clean("inscriptions", {"vessel_id": v.id, "transcription": vals[12],
                                                              "source": d["source"], "level": d["level"],
                                                              "evidence_grade": "B"}), status="pending", created_by=u.id)
                    db.add(i)
                    db.flush()
                    save_version(db, "inscriptions", i, TABLES["inscriptions"][1], u, "import")
                    retrieval.sync_inscription(db, i)
            else:
                d = dict(zip(["district", "speaker_code", "word", "ipa", "meaning", "source", "level"], vals[:7]))
                if not d["district"] or not d["word"]:
                    raise HTTPException(400, "区县与词条必填")
                e = DialectEntry(**_clean("dialect_entries", d), status="pending", created_by=u.id)
                db.add(e)
                db.flush()
                save_version(db, "dialect_entries", e, TABLES["dialect_entries"][1], u, "import")
                retrieval.sync_dialect(db, e)
            ok += 1
        except HTTPException as ex:
            errors.append({"line": line, "error": ex.detail})
    audit(db, u, "数据批量录入", kind, {"ok": ok, "failed": len(errors)}, client_ip(request))
    db.commit()
    return {"ok": ok, "errors": errors}


# ---------------- 裁断队列（RES-4） ----------------
def _adj(a: Adjudication, names) -> dict:
    return {"id": a.id, "title": a.title, "candidate": a.candidate, "question": a.question, "evidence": a.evidence,
            "confidence": a.confidence, "image_url": f"/api/files/{a.image_path}" if a.image_path else "",
            "status": a.status, "decision": a.decision, "revised_text": a.revised_text, "reason": a.reason,
            "draft": a.draft, "researcher": names.get(a.researcher_id, ""), "submitted_by": names.get(a.submitted_by, ""),
            "created_at": a.created_at.isoformat(timespec="minutes"),
            "decided_at": a.decided_at.isoformat(timespec="minutes") if a.decided_at else ""}


@router.get("/adjudications")
def adjudications(status: str = "pending", u: User = Depends(res_or_admin), db: Session = Depends(get_db)):
    rows = db.scalars(select(Adjudication).where(Adjudication.status == status).order_by(Adjudication.id.desc())).all()
    names = {x.id: x.name for x in db.scalars(select(User))}
    return [_adj(a, names) for a in rows]


class DecisionIn(BaseModel):
    decision: str  # adopt / revise / reject / draft
    revised_text: str = ""
    reason: str = ""


@router.post("/adjudications/{aid}")
def decide(aid: int, body: DecisionIn, request: Request, u: User = Depends(researcher), db: Session = Depends(get_db)):
    a = db.get(Adjudication, aid) or _404()
    if a.status != "pending":
        raise HTTPException(400, "该条已裁断")
    if body.decision == "draft":
        a.draft = {"revised_text": body.revised_text, "reason": body.reason}
        db.commit()
        return {"ok": True, "draft": True}
    if body.decision not in ("adopt", "revise", "reject"):
        raise HTTPException(400, "裁断方式无效")
    if not body.reason.strip():
        raise HTTPException(400, "裁断必须填写理由")
    if body.decision == "revise" and not body.revised_text.strip():
        raise HTTPException(400, "修订须填写修订结果")
    a.decision, a.revised_text, a.reason = body.decision, body.revised_text, body.reason
    a.status, a.researcher_id, a.decided_at = "done", u.id, datetime.now()
    retrieval.sync_adjudication(db, a)  # 采纳/修订结果进入检索（B 级）
    label = {"adopt": "采纳", "revise": "修订", "reject": "否决"}[body.decision]
    audit(db, u, "专家裁断", a.title, {"decision": label, "reason": body.reason}, client_ip(request))
    if a.submitted_by:
        db.add(Notification(user_id=a.submitted_by, kind="adjudication", title=f"“{a.title}”已裁断：{label}", body=body.reason[:200]))
    db.commit()
    return {"ok": True}


# ---------------- 任务审核（RES-6） ----------------
@router.get("/submissions")
def submissions(status: str = "pending", u: User = Depends(res_or_admin), db: Session = Depends(get_db)):
    rows = db.scalars(select(TaskSubmission).where(TaskSubmission.status == status).order_by(TaskSubmission.id.desc()).limit(200)).all()
    names = {x.id: x.name for x in db.scalars(select(User))}
    tpls = {t.id: t for t in db.scalars(select(TaskTemplate))}
    return [{"id": s.id, "template": tpls[s.template_id].name, "template_code": tpls[s.template_id].code,
             "user": names.get(s.user_id, ""), "target_title": s.target_title, "target_id": s.target_id,
             "content": s.content, "image_url": f"/api/files/{s.image_path}" if s.image_path else "",
             "status": s.status, "reward": s.reward, "reject_reason": s.reject_reason,
             "reviewer": names.get(s.reviewer_id, ""), "created_at": s.created_at.isoformat(timespec="minutes")} for s in rows]


@router.post("/submissions/{sid}/review")
def review_submission(sid: int, body: ReviewIn, request: Request, u: User = Depends(res_or_admin),
                      db: Session = Depends(get_db)):
    s = db.get(TaskSubmission, sid) or _404()
    if s.status != "pending":
        raise HTTPException(400, "该提交已审核")
    if s.user_id == u.id:
        raise HTTPException(403, "不能审核本人提交")
    tpl = db.get(TaskTemplate, s.template_id)
    if body.approve:
        reward = tpl.reward
        rule_key = {"correction": "reward_correction", "dialect": "reward_dialect"}.get(tpl.code)
        if rule_key:
            reward = points.rule(db, rule_key)
        s.status, s.reward = "approved", reward
        if reward:
            with points.user_lock(s.user_id):
                points.add_entry(db, s.user_id, "TASK_REWARD", reward, reason=f"{tpl.name}通过", ref_type="submission",
                                 ref_id=str(s.id), operator_id=u.id)
        db.add(Notification(user_id=s.user_id, kind="task", title=f"{tpl.name}已通过 +{reward}积分",
                            body=f"“{s.target_title}”的提交已由研究员审核通过。"))
    else:
        if not body.reason.strip():
            raise HTTPException(400, "驳回须写原因")
        s.status, s.reject_reason = "rejected", body.reason
        db.add(Notification(user_id=s.user_id, kind="task", title=f"{tpl.name}被驳回", body=body.reason[:200]))
    s.reviewer_id, s.reviewed_at = u.id, datetime.now()
    audit(db, u, "任务审核", f"提交#{s.id}", {"approve": body.approve, "reason": body.reason}, client_ip(request))
    db.commit()
    return {"ok": True}


class BatchIn(BaseModel):
    template_code: str = "dialect"
    title: str
    description: str = ""
    items: list[int] = []


@router.post("/task-batches")
def publish_batch(body: BatchIn, request: Request, u: User = Depends(researcher), db: Session = Depends(get_db)):
    tpl = db.scalar(select(TaskTemplate).where(TaskTemplate.code == body.template_code)) or _404()
    items = body.items or [i for (i,) in db.execute(select(DialectEntry.id).where(DialectEntry.status == "approved").limit(20))]
    b = TaskBatch(template_id=tpl.id, title=body.title, description=body.description, items=items, created_by=u.id)
    db.add(b)
    audit(db, u, "发布任务批次", body.title, {"items": len(items)}, client_ip(request))
    db.commit()
    return {"id": b.id}


@router.get("/workload")
def workload(u: User = Depends(res_or_admin), db: Session = Depends(get_db)):
    from ..models import AuditLog
    out = []
    for r in db.scalars(select(User).where(User.role == "researcher")):
        cnt = lambda act: db.scalar(select(func.count()).select_from(AuditLog).where(AuditLog.actor_id == r.id, AuditLog.action == act)) or 0
        out.append({"name": r.name, "entry": cnt("数据录入") + cnt("数据批量录入"), "review": cnt("数据审核"),
                    "adjudication": cnt("专家裁断"), "task_review": cnt("任务审核")})
    return sorted(out, key=lambda x: -(x["entry"] + x["review"] + x["adjudication"]))


def _404():
    raise HTTPException(404, "不存在")
